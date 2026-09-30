"""Turn bank/credit-card statements (CSV, Excel, PDF) into transaction dicts.

Every parser returns a list of {"date": "YYYY-MM-DD", "description": str,
"amount": float, "hint": str|None}. Amounts are signed: negative = money left
the account (a purchase, a bill), positive = money came in (paycheck, refund).
"""

import csv
import io
import re
from datetime import date, datetime

DATE_FORMATS = [
    "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y", "%m-%d-%Y", "%m-%d-%y",
    "%Y/%m/%d", "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d-%b-%Y", "%b %d %Y",
]


def parse_date(value: str, default_year: int | None = None) -> str | None:
    value = (value or "").strip()
    if not value:
        return None
    value = value.split("T")[0].split(" 00:00")[0]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    # Statement rows often omit the year: "09/14" or "Sep 14"
    if default_year:
        for fmt in ("%m/%d", "%m-%d", "%b %d", "%B %d"):
            try:
                d = datetime.strptime(value, fmt)
                return date(default_year, d.month, d.day).isoformat()
            except ValueError:
                continue
    return None


def parse_amount(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    negative = s.startswith("(") and s.endswith(")") or s.startswith("-") or s.endswith("-")
    if s.upper().endswith("CR"):
        negative = False
    digits = re.sub(r"[^0-9.]", "", s)
    if not digits or digits == ".":
        return None
    try:
        n = float(digits)
    except ValueError:
        return None
    return -n if negative else n


# ---------------------------------------------------------------- CSV / Excel

def _find(headers: list[str], *candidates: str) -> int | None:
    low = [h.lower().strip() for h in headers]
    for c in candidates:
        for i, h in enumerate(low):
            if h == c:
                return i
    for c in candidates:
        for i, h in enumerate(low):
            if c in h:
                return i
    return None


def parse_table(table: list[list], sign: str = "auto") -> list[dict]:
    """Parse rows from a spreadsheet-like statement. The header row is located
    automatically (some banks put a few lines of account info above it)."""
    header_idx = None
    for i, row in enumerate(table[:30]):
        cells = [str(c or "").lower() for c in row]
        if any("date" in c for c in cells) and any(
            k in c for c in cells for k in ("amount", "debit", "withdrawal", "credit", "deposit")
        ):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(
            "Couldn't find a header row with a date and an amount column. "
            "Export the statement as CSV from your bank's website, or upload the PDF."
        )
    headers = [str(c or "") for c in table[header_idx]]
    c_date = _find(headers, "transaction date", "trans. date", "date", "posted date", "posting date")
    c_desc = _find(headers, "description", "merchant", "payee", "name", "details", "memo", "transaction")
    c_amt = _find(headers, "amount")
    c_debit = _find(headers, "debit", "withdrawal", "withdrawals", "money out", "charges")
    c_credit = _find(headers, "credit", "deposit", "deposits", "money in", "payments")
    c_type = _find(headers, "type", "transaction type")
    c_cat = _find(headers, "category")
    if c_date is None or c_desc is None or (c_amt is None and c_debit is None and c_credit is None):
        raise ValueError(f"Unrecognized columns: {headers}")
    if c_amt is not None:
        # "Amount" wins; a column like "Credit Card Payments" shouldn't be mistaken for it.
        c_debit = c_credit = None

    out = []
    for row in table[header_idx + 1:]:
        row = list(row) + [None] * (len(headers) - len(row))
        d = parse_date(str(row[c_date] or ""))
        if not d:
            continue
        desc = " ".join(str(row[c_desc] or "").split())
        if c_amt is not None:
            amt = parse_amount(row[c_amt])
            if amt is None:
                continue
            if c_type is not None:
                t = str(row[c_type] or "").lower()
                if t in ("debit", "dr", "withdrawal", "sale", "purchase"):
                    amt = -abs(amt)
                elif t in ("credit", "cr", "deposit"):
                    amt = abs(amt)
        else:
            debit = parse_amount(row[c_debit]) if c_debit is not None else None
            credit = parse_amount(row[c_credit]) if c_credit is not None else None
            if debit:
                amt = -abs(debit)
            elif credit:
                amt = abs(credit)
            else:
                continue
        hint = str(row[c_cat]).strip() if c_cat is not None and row[c_cat] else None
        out.append({"date": d, "description": desc or "(no description)", "amount": amt, "hint": hint})

    return apply_sign(out, sign, single_column=c_amt is not None and c_type is None)


def apply_sign(txns: list[dict], sign: str, single_column: bool = True) -> list[dict]:
    """Normalize so purchases are negative. Some card issuers export purchases
    as positive numbers; `sign` = "auto" | "purchases_positive" | "purchases_negative"."""
    flip = sign == "purchases_positive"
    if sign == "auto" and single_column and txns:
        positives = sum(1 for t in txns if t["amount"] > 0)
        flip = positives / len(txns) > 0.7  # mostly-positive = card export with purchases as +
    if flip:
        for t in txns:
            t["amount"] = -t["amount"]
    return txns


def parse_csv(data: bytes, sign: str = "auto") -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    table = list(csv.reader(io.StringIO(text), dialect))
    return parse_table(table, sign)


def parse_xlsx(data: bytes, sign: str = "auto") -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    table = []
    for row in ws.iter_rows(values_only=True):
        table.append([c.strftime("%Y-%m-%d") if isinstance(c, (datetime, date)) else c for c in row])
    return parse_table(table, sign)


# ---------------------------------------------------------------- PDF

def pdf_text(data: bytes) -> list[str]:
    """Text of each page. Scanned (image-only) PDFs come back empty."""
    import pdfplumber

    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return pages


_YEAR_RX = re.compile(r"\b(20\d{2})\b")
_LINE_RX = re.compile(
    r"^\s*(?P<date>\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?|[A-Z][a-z]{2}\s+\d{1,2})\s+"
    r"(?:(?:\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?)\s+)?"  # optional second (post) date
    r"(?P<desc>.+?)\s+"
    r"(?P<amt>\(?-?\$?\s?[\d,]+\.\d{2}\)?(?:\s?CR|-)?)"
    r"(?:\s+\$?-?[\d,]+\.\d{2})?\s*$"  # optional running balance
)


def guess_statement_year(text: str) -> int:
    years = [int(y) for y in _YEAR_RX.findall(text) if 2000 <= int(y) <= date.today().year + 1]
    return max(set(years), key=years.count) if years else date.today().year


def parse_pdf_regex(pages: list[str], sign: str = "auto") -> list[dict]:
    """Line-by-line fallback used when the AI isn't configured. Works for most
    statements that print one transaction per line; the AI path is more robust."""
    text = "\n".join(pages)
    year = guess_statement_year(text)
    out = []
    for line in text.splitlines():
        m = _LINE_RX.match(line)
        if not m:
            continue
        d = parse_date(m["date"], default_year=year)
        amt = parse_amount(m["amt"])
        if not d or amt is None:
            continue
        desc = " ".join(m["desc"].split())
        if re.search(r"balance|total|minimum payment|payment due", desc, re.I):
            continue
        out.append({"date": d, "description": desc, "amount": amt, "hint": None})
    return apply_sign(out, sign)
