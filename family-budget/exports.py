"""PDF summary, Excel/CSV export and (optional) direct Google Sheets export."""

import csv
import io
import os

import budget


def _money(v: float) -> str:
    return f"-${abs(v):,.0f}" if v < 0 else f"${v:,.0f}"


# ---------------------------------------------------------------- PDF

def pdf_summary(month: str, family: str = "Family") -> bytes:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    green, red, ink, muted = colors.HexColor("#2e9e6a"), colors.HexColor("#d23f3f"), colors.HexColor("#1f1d1a"), colors.HexColor("#77736c")
    line, band = colors.HexColor("#e6e3de"), colors.HexColor("#f5f4f2")

    b = budget.month_budget(month)
    t = b["totals"]
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], alignment=0, fontSize=20, textColor=ink, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=muted, fontSize=10)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=ink, fontSize=13, spaceBefore=14, spaceAfter=6)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, textColor=ink, leading=10)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            title=f"{family} budget — {b['label']}")
    story = [Paragraph(f"{family} Budget Summary", h1), Paragraph(b["label"], sub), Spacer(1, 12)]

    # KPI strip
    kpis = [
        ("Income", _money(t["income_actual"]), f"of {_money(t['income_planned'])} planned"),
        ("Expenses", _money(t["expense_actual"]), f"of {_money(t['expense_planned'])} planned"),
        ("Net saved", _money(t["net_actual"]), f"{t['savings_rate']}% savings rate" if t["savings_rate"] is not None else ""),
        ("Left to budget", _money(t["left_to_budget"]), "planned income minus expenses"),
    ]
    k_label = ParagraphStyle("kl", parent=styles["Normal"], fontSize=8, leading=11, textColor=muted)
    k_value = ParagraphStyle("kv", parent=styles["Normal"], fontSize=15, leading=19, textColor=ink, fontName="Helvetica-Bold")
    k_sub = ParagraphStyle("ks", parent=styles["Normal"], fontSize=7, leading=9, textColor=muted)
    kpi_table = Table(
        [[[Paragraph(k, k_label), Paragraph(v, k_value), Paragraph(s, k_sub)] for k, v, s in kpis]],
        colWidths=[doc.width / 4] * 4, rowHeights=[0.8 * inch],
    )
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), band), ("BOX", (0, 0), (-1, -1), 0.5, line),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, line), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(kpi_table)

    # Spending by category chart
    spend = sorted(
        [(c["name"], c["actual"]) for g in b["groups"] if g["key"] in budget.EXPENSE_GROUPS
         for c in g["categories"] if c["actual"] > 0],
        key=lambda x: -x[1],
    )[:10]
    if spend:
        story.append(Paragraph("Where the money went", h2))
        d = Drawing(doc.width, 18 * len(spend) + 20)
        chart = HorizontalBarChart()
        chart.x, chart.y, chart.width, chart.height = 130, 10, doc.width - 170, 18 * len(spend)
        chart.data = [[v for _, v in reversed(spend)]]
        chart.categoryAxis.categoryNames = [n for n, _ in reversed(spend)]
        chart.categoryAxis.labels.fontSize = 8
        chart.categoryAxis.strokeColor = line
        chart.valueAxis.valueMin = 0
        chart.valueAxis.labels.fontSize = 7
        chart.valueAxis.strokeColor = line
        chart.valueAxis.labelTextFormat = lambda v: f"${v:,.0f}"
        chart.bars[0].fillColor = colors.HexColor("#4a7fd4")
        chart.bars[0].strokeColor = None
        chart.barLabelFormat = lambda v: f"${v:,.0f}"
        chart.barLabels.fontSize = 7
        chart.barLabels.boxAnchor = "w"
        chart.barLabels.dx = 3
        d.add(chart)
        story.append(d)

    # Budget tables
    for g in b["groups"]:
        rows = [c for c in g["categories"] if c["planned"] or c["actual"]]
        if not rows:
            continue
        data = [[g["label"], "Planned", "Actual", "Remaining"]]
        style = [
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("TEXTCOLOR", (1, 0), (-1, 0), muted), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, line), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
        for i, c in enumerate(rows, start=1):
            data.append([c["name"], _money(c["planned"]), _money(c["actual"]),
                         _money(c["remaining"]) if c["planned"] else ""])
            if c["planned"] and g["key"] != "income" and c["remaining"] < 0:
                style.append(("TEXTCOLOR", (3, i), (3, i), red))
            elif c["planned"] and g["key"] != "income" and c["remaining"] >= 0:
                style.append(("TEXTCOLOR", (3, i), (3, i), green))
        data.append([f"Total {g['label']}", _money(g["planned"]), _money(g["actual"]), _money(g["remaining"])])
        n = len(data) - 1
        style += [("FONTNAME", (0, n), (-1, n), "Helvetica-Bold"), ("BACKGROUND", (0, n), (-1, n), band)]
        if g["key"] == "transfer":
            data[0][0] += " (excluded from totals)"
        tbl = Table(data, colWidths=[doc.width * 0.46] + [doc.width * 0.18] * 3, repeatRows=1)
        tbl.setStyle(TableStyle(style))
        story += [Spacer(1, 10), KeepTogether(tbl)]

    # Transaction appendix
    txns = budget.transactions(month)
    if txns:
        story += [PageBreak(), Paragraph(f"All transactions ({len(txns)})", h2)]
        data = [["Date", "Description", "Category", "Amount"]]
        for x in txns:
            data.append([x["date"], Paragraph(x["description"][:70], small), x["category"] or "Uncategorized",
                         f"{x['amount']:,.2f}"])
        tbl = Table(data, colWidths=[0.8 * inch, doc.width - 3.4 * inch, 1.6 * inch, 1.0 * inch], repeatRows=1)
        tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (3, 0), (3, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.3, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(tbl)

    def footer(canvas, _doc):
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        canvas.drawString(doc.leftMargin, 0.35 * inch, f"{family} budget · {b['label']}")
        canvas.drawRightString(letter[0] - doc.rightMargin, 0.35 * inch, f"Page {canvas.getPageNumber()}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------- tabular exports

def summary_rows(month: str) -> list[list]:
    b = budget.month_budget(month)
    rows = [["Group", "Category", "Planned", "Actual", "Remaining", "Transactions"]]
    for g in b["groups"]:
        for c in g["categories"]:
            if c["planned"] or c["actual"]:
                rows.append([g["label"], c["name"], c["planned"], c["actual"], c["remaining"], c["count"]])
        rows.append([g["label"], "TOTAL", g["planned"], g["actual"], g["remaining"], ""])
    t = b["totals"]
    rows += [[], ["", "Total income", t["income_planned"], t["income_actual"], t["income_remaining"], ""],
             ["", "Total expenses", t["expense_planned"], t["expense_actual"], t["expense_remaining"], ""],
             ["", "Left to budget", t["left_to_budget"], "", "", ""],
             ["", "Net (actual)", "", t["net_actual"], "", ""]]
    return rows


def transaction_rows(month: str | None) -> list[list]:
    rows = [["Date", "Description", "Merchant", "Category", "Group", "Amount", "Account", "Notes"]]
    for x in budget.transactions(month):
        rows.append([x["date"], x["description"], x["merchant"] or "", x["category"] or "Uncategorized",
                     x["grp"] or "", x["amount"], x["account"] or "", x["notes"] or ""])
    return rows


def trend_rows(month: str) -> list[list]:
    rows = [["Month", "Income", "Expenses", "Net"]]
    for r in budget.trend(month, 12):
        rows.append([budget.month_label(r["month"]), r["income"], r["expenses"], round(r["income"] - r["expenses"], 2)])
    return rows


def xlsx(month: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    sheets = [("Summary", summary_rows(month)), ("Transactions", transaction_rows(month)),
              ("12-Month Trend", trend_rows(month))]
    for i, (title, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EDEAE5")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '"$"#,##0.00;[Red]-"$"#,##0.00'
                if row[1].value in ("TOTAL", "Left to budget", "Net (actual)"):
                    cell.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_bytes(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------- Google Sheets

def google_sheets_configured() -> bool:
    path = os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE")
    return bool(path and os.path.exists(path))


def export_google_sheet(month: str, share_with: str | None, family: str) -> str:
    """Create a Google Sheet with Summary/Transactions/Trend tabs and share it.
    Requires GOOGLE_SERVICE_ACCOUNT_FILE (a service-account JSON key) and `gspread`."""
    import gspread

    gc = gspread.service_account(filename=os.environ["GOOGLE_SERVICE_ACCOUNT_FILE"])
    sh = gc.create(f"{family} Budget — {budget.month_label(month)}")
    tabs = [("Summary", summary_rows(month)), ("Transactions", transaction_rows(month)),
            ("12-Month Trend", trend_rows(month))]
    for i, (title, rows) in enumerate(tabs):
        rows = [[("" if v is None else v) for v in r] for r in rows]
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        ws = sh.sheet1 if i == 0 else sh.add_worksheet(title=title, rows=len(rows) + 10, cols=width)
        if i == 0:
            ws.update_title(title)
        ws.update(rows, "A1")
        ws.format("1:1", {"textFormat": {"bold": True}})
    if share_with:
        sh.share(share_with, perm_type="user", role="writer")
    return sh.url
